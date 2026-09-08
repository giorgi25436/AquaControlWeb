from functools import wraps
import os
import sqlite3
import secrets
from datetime import datetime
from datetime import date, timedelta
from io import BytesIO

from flask import Flask, flash, redirect, render_template, request, session, url_for
from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph
from werkzeug.security import check_password_hash, generate_password_hash
from database import audit
from cycles import calculate_cycle
from database import active_worker_count, attendance_counts, get_connection, init_db

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config["DEBUG"] = os.environ.get("FLASK_DEBUG", "").lower() in {"1", "true", "yes"}
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=8)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.environ.get("SESSION_COOKIE_SECURE", "").lower() in {"1", "true", "yes"}
init_db()
with get_connection() as connection:
    if connection.execute("SELECT COUNT(*) FROM usuarios").fetchone()[0] == 0:
        connection.execute(
            """
            INSERT INTO usuarios (nombre, usuario, password_hash, rol, estado, fecha_creacion)
            VALUES (?, ?, ?, 'Administrador', 'Activo', ?)
            """,
            (
                "Administrador inicial",
                "admin",
                generate_password_hash(os.environ.get("INITIAL_ADMIN_PASSWORD", "Aqua2026!")),
                datetime.now().isoformat(timespec="seconds"),
            ),
        )


def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if "username" not in session:
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped_view


def role_required(*roles):
    def decorator(view):
        @wraps(view)
        def wrapped_view(*args, **kwargs):
            if "user_id" not in session:
                return redirect(url_for("login"))
            if session.get("rol") not in roles:
                flash("No tienes permisos para realizar esta acción.", "error")
                return redirect(url_for("dashboard"))
            return view(*args, **kwargs)
        return wrapped_view
    return decorator


def current_user_id():
    return session.get("user_id")


@app.route("/")
def index():
    return redirect(url_for("login"))


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        with get_connection() as connection:
            user = connection.execute(
                "SELECT * FROM usuarios WHERE usuario = ?", (username,)
            ).fetchone()
        if user and user["estado"] == "Activo" and check_password_hash(user["password_hash"], password):
            session.clear()
            session.permanent = True
            session.update({
                "user_id": user["id"],
                "nombre": user["nombre"],
                "username": user["usuario"],
                "rol": user["rol"],
            })
            audit(user["id"], "Inicio de sesión")
            return redirect(url_for("dashboard"))

        flash("Usuario o contraseña incorrectos", "error")

    return render_template("login.html")


@app.route("/dashboard")
@login_required
def dashboard():
    today = date.today()
    cycles = get_cycle_data(today)
    resting_groups = tuple(item["name"] for item in cycles if item["is_resting"])
    today_counts = attendance_counts(today.isoformat(), resting_groups)
    alerts = cycle_alerts(cycles)
    exit_candidates = [item for item in cycles if item["next_exit"] != date.max]
    return_candidates = [item for item in cycles if item["next_return"] != date.max]
    return render_template(
        "dashboard.html",
        username=session["username"],
        user_name=session["nombre"],
        user_role=session["rol"],
        active_worker_count=active_worker_count(),
        present_count=today_counts.get("Presente", 0),
        absent_count=today_counts.get("Ausente", 0),
        cycles=cycles,
        alerts=alerts,
        next_exit=min(exit_candidates, key=lambda item: item["next_exit"]) if exit_candidates else None,
        next_return=min(return_candidates, key=lambda item: item["next_return"]) if return_candidates else None,
        resting_group_count=sum(item["is_resting"] for item in cycles),
    )


@app.route("/trabajadores", methods=["GET", "POST"])
@role_required("Administrador", "Supervisor", "Consulta")
def trabajadores():
    if request.method == "POST":
        if session["rol"] not in {"Administrador"}:
            flash("No tienes permisos para realizar esta acción.", "error")
            return redirect(url_for("trabajadores"))
        worker_id = request.form.get("worker_id", "").strip()
        data = {
            "cedula": request.form.get("cedula", "").strip(),
            "nombres": request.form.get("nombres", "").strip(),
            "apellidos": request.form.get("apellidos", "").strip(),
            "cargo": request.form.get("cargo", "").strip(),
            "telefono": request.form.get("telefono", "").strip(),
            "direccion": request.form.get("direccion", "").strip(),
            "fecha_ingreso": request.form.get("fecha_ingreso", "").strip(),
            "grupo": request.form.get("grupo", "").strip(),
            "estado": request.form.get("estado", "").strip(),
            "observaciones": request.form.get("observaciones", "").strip(),
        }
        required_fields = data.copy()
        required_fields.pop("observaciones")
        if not all(required_fields.values()) or data["grupo"] not in {"A", "B", "C"} or data["estado"] not in {"Activo", "Inactivo"}:
            flash("Completa todos los campos obligatorios con valores válidos.", "error")
        else:
            try:
                with get_connection() as connection:
                    if worker_id:
                        connection.execute(
                            """
                            UPDATE trabajadores SET cedula=:cedula, nombres=:nombres,
                            apellidos=:apellidos, cargo=:cargo, telefono=:telefono,
                            direccion=:direccion, fecha_ingreso=:fecha_ingreso, grupo=:grupo,
                            estado=:estado, observaciones=:observaciones WHERE id=:id
                            """,
                            {**data, "id": worker_id},
                        )
                        flash("Trabajador actualizado correctamente.", "success")
                        audit(current_user_id(), "Edición de trabajador", data["cedula"], connection)
                    else:
                        connection.execute(
                            """
                            INSERT INTO trabajadores
                            (cedula, nombres, apellidos, cargo, telefono, direccion,
                             fecha_ingreso, grupo, estado, observaciones)
                            VALUES (:cedula, :nombres, :apellidos, :cargo, :telefono,
                                    :direccion, :fecha_ingreso, :grupo, :estado, :observaciones)
                            """,
                            data,
                        )
                        flash("Trabajador registrado correctamente.", "success")
                        audit(current_user_id(), "Registro de trabajador", data["cedula"], connection)
                return redirect(url_for("trabajadores"))
            except sqlite3.IntegrityError as error:
                if "UNIQUE constraint failed" in str(error):
                    flash("Ya existe un trabajador con esa cédula.", "error")
                else:
                    raise

    search = request.args.get("buscar", "").strip()
    group = request.args.get("grupo", "").strip()
    query = "SELECT * FROM trabajadores WHERE 1=1"
    params = []
    if search:
        query += " AND (cedula LIKE ? OR nombres LIKE ? OR apellidos LIKE ?)"
        params.extend([f"%{search}%"] * 3)
    if group in {"A", "B", "C"}:
        query += " AND grupo = ?"
        params.append(group)
    query += " ORDER BY id DESC"
    with get_connection() as connection:
        workers = connection.execute(query, params).fetchall()
        worker_to_edit = None
        edit_id = request.args.get("editar", "").strip()
        if edit_id:
            worker_to_edit = connection.execute(
                "SELECT * FROM trabajadores WHERE id = ?", (edit_id,)
            ).fetchone()
    return render_template(
        "trabajadores.html",
        workers=workers,
        worker=worker_to_edit,
        search=search,
        selected_group=group,
    )


@app.post("/trabajadores/<int:worker_id>/eliminar")
@role_required("Administrador")
def eliminar_trabajador(worker_id):
    with get_connection() as connection:
        connection.execute("DELETE FROM trabajadores WHERE id = ?", (worker_id,))
    flash("Trabajador eliminado correctamente.", "success")
    audit(current_user_id(), "Eliminación de trabajador", str(worker_id))
    return redirect(url_for("trabajadores"))


@app.route("/asistencia", methods=["GET", "POST"])
@role_required("Administrador", "Supervisor")
def asistencia():
    today = datetime.now().strftime("%Y-%m-%d")
    now = datetime.now().strftime("%H:%M")
    cedula = request.values.get("cedula", "").strip()
    worker = None
    today_attendance = None

    if request.method == "POST":
        action = request.form.get("action", "").strip()
        with get_connection() as connection:
            worker = connection.execute(
                "SELECT * FROM trabajadores WHERE cedula = ?", (cedula,)
            ).fetchone()
            if worker is None:
                flash("El trabajador no está registrado.", "error")
            else:
                worker_cycle = next(
                    item for item in get_cycle_data(date.today())
                    if item["name"] == worker["grupo"]
                )
                if worker_cycle["is_resting"]:
                    flash("DESCANSO PROGRAMADO: no es necesario registrar asistencia.", "error")
                    return redirect(url_for("asistencia", cedula=cedula))
                today_attendance = connection.execute(
                    "SELECT * FROM asistencias WHERE trabajador_id = ? AND fecha = ?",
                    (worker["id"], today),
                ).fetchone()
                if action == "entrada":
                    if today_attendance:
                        flash("La entrada de este trabajador ya fue registrada hoy.", "error")
                    else:
                        connection.execute(
                            """
                            INSERT INTO asistencias
                            (trabajador_id, cedula, fecha, hora_entrada, estado, observacion)
                            VALUES (?, ?, ?, ?, 'Presente', '')
                            """,
                            (worker["id"], worker["cedula"], today, now),
                        )
                        flash("Entrada registrada correctamente.", "success")
                        audit(current_user_id(), "Registro de asistencia", f"Entrada {worker['cedula']}", connection)
                elif action == "salida":
                    if not today_attendance:
                        flash("Primero debes registrar la entrada del trabajador.", "error")
                    elif today_attendance["hora_salida"]:
                        flash("La salida de este trabajador ya fue registrada hoy.", "error")
                    else:
                        connection.execute(
                            "UPDATE asistencias SET hora_salida = ? WHERE id = ?",
                            (now, today_attendance["id"]),
                        )
                        flash("Salida registrada correctamente.", "success")
                        audit(current_user_id(), "Registro de asistencia", f"Salida {worker['cedula']}", connection)
                elif action in {"permiso", "vacaciones"}:
                    if today_attendance:
                        flash("Este trabajador ya tiene una asistencia registrada hoy.", "error")
                    else:
                        state = "Permiso" if action == "permiso" else "Vacaciones"
                        observation = request.form.get("observacion", "").strip()
                        connection.execute(
                            """
                            INSERT INTO asistencias
                            (trabajador_id, cedula, fecha, estado, observacion)
                            VALUES (?, ?, ?, ?, ?)
                            """,
                            (worker["id"], worker["cedula"], today, state, observation),
                        )
                        flash(f"{state} registrada correctamente.", "success")
                        audit(current_user_id(), "Registro de asistencia", f"{state} {worker['cedula']}", connection)
                return redirect(url_for("asistencia", cedula=cedula))

    with get_connection() as connection:
        if cedula:
            worker = connection.execute(
                "SELECT * FROM trabajadores WHERE cedula = ?", (cedula,)
            ).fetchone()
            if worker:
                today_attendance = connection.execute(
                    "SELECT * FROM asistencias WHERE trabajador_id = ? AND fecha = ?",
                    (worker["id"], today),
                ).fetchone()
        worker_cycle = None
        if worker:
            worker_cycle = next(
                (item for item in get_cycle_data(date.today()) if item["name"] == worker["grupo"]),
                None,
            )
        attendance_query = """
            SELECT a.*, t.nombres, t.apellidos, t.grupo
            FROM asistencias a
            JOIN trabajadores t ON t.id = a.trabajador_id
            WHERE a.fecha = ?
        """
        attendance_params = [today]
        table_search = request.args.get("buscar", "").strip()
        if table_search:
            attendance_query += " AND (a.cedula LIKE ? OR t.nombres LIKE ? OR t.apellidos LIKE ?)"
            attendance_params.extend([f"%{table_search}%"] * 3)
        attendance_query += " ORDER BY a.id DESC"
        todays_attendance = connection.execute(
            attendance_query, attendance_params
        ).fetchall()
    return render_template(
        "asistencia.html",
        worker=worker,
        today_attendance=today_attendance,
        todays_attendance=todays_attendance,
        cedula=cedula,
        table_search=request.args.get("buscar", "").strip(),
        today=today,
        worker_cycle=worker_cycle,
    )


def get_cycle_data(today=None):
    today = today or date.today()
    with get_connection() as connection:
        rows = connection.execute(
            "SELECT grupo, fecha_inicio, modo_estado, fecha_inicio_manual, fecha_regreso_manual FROM ciclos ORDER BY grupo"
        ).fetchall()
    return [
        calculate_cycle(
            row["grupo"],
            date.fromisoformat(row["fecha_inicio"]),
            today,
            row["modo_estado"],
            date.fromisoformat(row["fecha_inicio_manual"]) if row["fecha_inicio_manual"] else None,
            date.fromisoformat(row["fecha_regreso_manual"]) if row["fecha_regreso_manual"] else None,
        )
        for row in rows
    ]


def cycle_alerts(cycles):
    alerts = []
    for item in cycles:
        days = item["days_to_transition"]
        if days in {1, 2, 3}:
            if item["is_resting"]:
                text = f"El Grupo {item['name']} regresará " + ("mañana." if days == 1 else f"en {days} días.")
            else:
                text = f"El Grupo {item['name']} saldrá de descanso " + ("hoy." if days == 1 else f"en {days} días.")
            alerts.append(text)
    return alerts


@app.route("/ciclos", methods=["GET", "POST"])
@role_required("Administrador")
def ciclos():
    if request.method == "POST":
        for group in ("A", "B", "C"):
            start = request.form.get(f"fecha_{group}", "").strip()
            mode = request.form.get(f"modo_{group}", "Automatico").strip()
            manual_start = request.form.get(f"manual_inicio_{group}", "").strip() or None
            manual_return = request.form.get(f"manual_regreso_{group}", "").strip() or None
            try:
                date.fromisoformat(start)
                if mode not in {"Automatico", "Forzar LABORANDO", "Forzar EN DESCANSO"}:
                    raise ValueError
                if mode != "Automatico" and not manual_start:
                    raise ValueError
                manual_start_date = date.fromisoformat(manual_start) if manual_start else None
                manual_return_date = date.fromisoformat(manual_return) if manual_return else None
                if mode == "Forzar EN DESCANSO" and manual_return_date and manual_return_date <= manual_start_date:
                    raise ValueError
            except ValueError:
                flash(f"La fecha del Grupo {group} no es válida.", "error")
                return redirect(url_for("ciclos"))
            with get_connection() as connection:
                connection.execute(
                    """
                    UPDATE ciclos SET fecha_inicio = ?, dias_trabajo = 22, dias_descanso = 8,
                    modo_estado = ?, fecha_inicio_manual = ?, fecha_regreso_manual = ?
                    WHERE grupo = ?
                    """,
                    (start, mode, manual_start, manual_return, group),
                )
        flash("Fechas de ciclo actualizadas correctamente.", "success")
        audit(current_user_id(), "Cambio de ciclos")
        return redirect(url_for("ciclos"))
    cycle_data = get_cycle_data()
    return render_template("ciclos.html", cycles=cycle_data, alerts=cycle_alerts(cycle_data))


@app.route("/calendario")
@role_required("Administrador", "Supervisor", "Consulta")
def calendario():
    today = date.today()
    try:
        year = int(request.args.get("year", today.year))
        month = int(request.args.get("month", today.month))
        selected_date = date.fromisoformat(request.args.get("fecha", today.isoformat()))
        if not 1 <= month <= 12:
            raise ValueError
    except ValueError:
        year, month, selected_date = today.year, today.month, today

    first_day = date(year, month, 1)
    next_month = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
    previous_month = date(year - (month == 1), 12 if month == 1 else month - 1, 1)
    grid_start = first_day - timedelta(days=first_day.weekday())
    grid_end = next_month + timedelta(days=(6 - next_month.weekday()))
    calendar_days = []
    current = grid_start
    while current < grid_end:
        day_cycles = get_cycle_data(current)
        previous_cycles = get_cycle_data(current - timedelta(days=1))
        day_data = []
        for item, previous in zip(day_cycles, previous_cycles):
            event = "regreso" if previous["is_resting"] and not item["is_resting"] else ""
            event = "salida" if not previous["is_resting"] and item["is_resting"] else event
            day_data.append({**item, "event": event})
        calendar_days.append({"date": current, "cycles": day_data, "is_current_month": current.month == month})
        current += timedelta(days=1)

    selected_cycles = get_cycle_data(selected_date)
    previous_selected_cycles = get_cycle_data(selected_date - timedelta(days=1))
    selected_cycles = [
        {
            **item,
            "event": (
                "regreso" if previous["is_resting"] and not item["is_resting"]
                else "salida" if not previous["is_resting"] and item["is_resting"]
                else ""
            ),
        }
        for item, previous in zip(selected_cycles, previous_selected_cycles)
    ]
    movements = []
    cursor = today
    for _ in range(90):
        current_cycles = get_cycle_data(cursor)
        previous_cycles = get_cycle_data(cursor - timedelta(days=1))
        for item, previous in zip(current_cycles, previous_cycles):
            if item["is_resting"] != previous["is_resting"]:
                movements.append({
                    "date": cursor,
                    "group": item["name"],
                    "type": "Salida de descanso" if item["is_resting"] else "Regreso a trabajar",
                    "is_exit": item["is_resting"],
                })
        cursor += timedelta(days=1)
    return render_template(
        "calendario.html",
        calendar_days=calendar_days,
        selected_date=selected_date,
        selected_cycles=selected_cycles,
        month_name={
            1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril", 5: "Mayo", 6: "Junio",
            7: "Julio", 8: "Agosto", 9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre",
        }[month] + f" {year}",
        year=year,
        month=month,
        previous_month=previous_month,
        next_month=next_month,
        movements=movements[:12],
        today=today,
    )


@app.route("/historial")
@role_required("Administrador", "Supervisor", "Consulta")
def historial():
    search = request.args.get("buscar", "").strip()
    date_from = request.args.get("fecha_desde", "").strip()
    date_to = request.args.get("fecha_hasta", "").strip()
    month = request.args.get("mes", "").strip()
    year = request.args.get("ano", "").strip()
    state = request.args.get("estado", "").strip()
    worker = None
    records = []
    summary = {key: 0 for key in ("Presente", "Ausente", "Permiso", "Vacaciones", "Descanso programado")}

    with get_connection() as connection:
        if search:
            worker = connection.execute(
                """
                SELECT * FROM trabajadores
                WHERE cedula LIKE ? OR nombres LIKE ? OR apellidos LIKE ?
                ORDER BY nombres, apellidos LIMIT 1
                """,
                (f"%{search}%", f"%{search}%", f"%{search}%"),
            ).fetchone()
        if worker:
            try:
                start = date.fromisoformat(date_from) if date_from else date.fromisoformat(worker["fecha_ingreso"])
                end = date.fromisoformat(date_to) if date_to else date.today()
                if month:
                    month_number = int(month)
                    start = date(int(year) if year else date.today().year, month_number, 1)
                    end = date(start.year + (month_number == 12), 1 if month_number == 12 else month_number + 1, 1) - timedelta(days=1)
                if year:
                    year_number = int(year)
                    selected_month = int(month) if month else 1
                    start = date(year_number, selected_month, 1)
                    if month:
                        next_year = year_number + (selected_month == 12)
                        next_month = 1 if selected_month == 12 else selected_month + 1
                        end = date(next_year, next_month, 1) - timedelta(days=1)
                    else:
                        end = date(year_number, 12, 31)
                end = min(end, date.today())
                if end < start:
                    raise ValueError
            except ValueError:
                flash("El rango de fechas seleccionado no es válido.", "error")
                start, end = date.today(), date.today()

            stored = connection.execute(
                """
                SELECT a.*, t.grupo FROM asistencias a
                JOIN trabajadores t ON t.id = a.trabajador_id
                WHERE a.trabajador_id = ? AND a.fecha BETWEEN ? AND ?
                ORDER BY a.fecha DESC
                """,
                (worker["id"], start.isoformat(), end.isoformat()),
            ).fetchall()
            stored_by_date = {row["fecha"]: row for row in stored}
            current = start
            while current <= end:
                cycle = next(item for item in get_cycle_data(current) if item["name"] == worker["grupo"])
                row = stored_by_date.get(current.isoformat())
                if cycle["is_resting"] and (not row or row["estado"] == "Ausente"):
                    item = {"fecha": current.isoformat(), "hora_entrada": None, "hora_salida": None, "estado": "Descanso programado", "grupo": worker["grupo"], "observacion": "Descanso según ciclo 22x8"}
                elif row:
                    item = dict(row)
                else:
                    current += timedelta(days=1)
                    continue
                if not state or item["estado"] == state:
                    records.append(item)
                    summary[item["estado"]] += 1
                current += timedelta(days=1)

    return render_template(
        "historial.html",
        worker=worker,
        records=records,
        summary=summary,
        search=search,
        date_from=date_from,
        date_to=date_to,
        month=month,
        year=year,
        state=state,
        weekday_name=weekday_name,
    )


def report_filters():
    return {
        "date_from": request.args.get("fecha_desde", "").strip(),
        "date_to": request.args.get("fecha_hasta", "").strip(),
        "month": request.args.get("mes", "").strip(),
        "year": request.args.get("ano", "").strip(),
        "worker": request.args.get("trabajador", "").strip(),
        "cedula": request.args.get("cedula", "").strip(),
        "group": request.args.get("grupo", "").strip(),
        "state": request.args.get("estado", "").strip(),
        "report_type": request.args.get("tipo", "diario").strip(),
    }


def build_report_data(filters):
    today = date.today()
    start = date.fromisoformat(filters["date_from"]) if filters["date_from"] else today
    end = date.fromisoformat(filters["date_to"]) if filters["date_to"] else today
    if filters["report_type"] == "semanal" and not filters["date_from"]:
        start = today - timedelta(days=today.weekday())
    elif filters["report_type"] == "mensual" and not filters["date_from"]:
        start = today.replace(day=1)
    if filters["month"]:
        month = int(filters["month"])
        start = date(int(filters["year"]) if filters["year"] else today.year, month, 1)
        end = date(start.year + (month == 12), 1 if month == 12 else month + 1, 1) - timedelta(days=1)
    elif filters["year"]:
        start = date(int(filters["year"]), 1, 1)
        end = date(int(filters["year"]), 12, 31)
    end = min(end, today)
    if end < start:
        raise ValueError("El período seleccionado no es válido.")

    with get_connection() as connection:
        worker_query = "SELECT * FROM trabajadores WHERE 1=1"
        worker_params = []
        if filters["worker"]:
            worker_query += " AND (nombres LIKE ? OR apellidos LIKE ?)"
            worker_params.extend([f"%{filters['worker']}%"] * 2)
        if filters["cedula"]:
            worker_query += " AND cedula LIKE ?"
            worker_params.append(f"%{filters['cedula']}%")
        if filters["group"] in {"A", "B", "C"}:
            worker_query += " AND grupo = ?"
            worker_params.append(filters["group"])
        workers = connection.execute(worker_query, worker_params).fetchall()
        actual = connection.execute(
            """
            SELECT a.*, t.nombres, t.apellidos, t.cargo, t.grupo
            FROM asistencias a JOIN trabajadores t ON t.id = a.trabajador_id
            WHERE a.fecha BETWEEN ? AND ?
            """,
            (start.isoformat(), end.isoformat()),
        ).fetchall()
    actual_by_worker_date = {(row["trabajador_id"], row["fecha"]): dict(row) for row in actual}
    rows = []
    current = start
    while current <= end:
        for worker in workers:
            row = actual_by_worker_date.get((worker["id"], current.isoformat()))
            cycle = next(item for item in get_cycle_data(current) if item["name"] == worker["grupo"])
            if cycle["is_resting"] and (not row or row["estado"] == "Ausente"):
                row = {
                    "fecha": current.isoformat(), "cedula": worker["cedula"],
                    "nombres": worker["nombres"], "apellidos": worker["apellidos"],
                    "cargo": worker["cargo"], "grupo": worker["grupo"],
                    "hora_entrada": None, "hora_salida": None,
                    "estado": "Descanso programado",
                    "observacion": "Descanso según ciclo 22x8",
                }
            elif not row:
                continue
            if filters["state"] and row["estado"] != filters["state"]:
                continue
            rows.append(row)
        current += timedelta(days=1)
    rows.sort(key=lambda row: row["fecha"], reverse=True)
    summary = {key: 0 for key in ("Presente", "Ausente", "Permiso", "Vacaciones", "Descanso programado")}
    for row in rows:
        summary[row["estado"]] += 1
    return rows, summary, workers, start, end


@app.route("/reportes")
@role_required("Administrador", "Supervisor", "Consulta")
def reportes():
    filters = report_filters()
    rows, summary, workers, start, end = [], {key: 0 for key in ("Presente", "Ausente", "Permiso", "Vacaciones", "Descanso programado")}, [], None, None
    error = None
    if request.args:
        try:
            rows, summary, workers, start, end = build_report_data(filters)
        except (ValueError, TypeError):
            error = "Los filtros del reporte no son válidos."
    return render_template("reportes.html", **filters, rows=rows, summary=summary, workers=workers, start=start, end=end, error=error)


def export_report(format_name):
    filters = report_filters()
    rows, summary, _, start, end = build_report_data(filters)
    title = f"Reporte {filters['report_type'].capitalize()}"
    if format_name == "excel":
        book = Workbook()
        sheet = book.active
        sheet.title = "Reporte"
        sheet.append(["AGCOMPANY"])
        sheet.append(["Gestión de Personal Camaronero"])
        sheet.append([title, f"Período: {start.strftime('%d/%m/%Y')} - {end.strftime('%d/%m/%Y')}"])
        sheet.append([])
        headers = ["Fecha", "Cédula", "Trabajador", "Cargo", "Grupo", "Hora de entrada", "Hora de salida", "Estado", "Observación"]
        sheet.append(headers)
        for row in rows:
            sheet.append([row["fecha"], row["cedula"], f"{row['nombres']} {row['apellidos']}", row["cargo"], row["grupo"], row["hora_entrada"] or "", row["hora_salida"] or "", row["estado"], row["observacion"] or ""])
        sheet.append([])
        sheet.append(["Totales"] + [f"{key}: {value}" for key, value in summary.items()])
        sheet["A1"].font = Font(bold=True, size=16, color="1264D6")
        for cell in sheet[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B1F3A")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(max(column, key=lambda cell: len(str(cell.value or ""))).value)) + 2, 12),  thirty := 30)
        output = BytesIO()
        book.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name="aqua_control_reporte.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(letter), rightMargin=25, leftMargin=25, topMargin=25, bottomMargin=25)
    styles = getSampleStyleSheet()
    story = [Paragraph("AGCOMPANY", styles["Title"]), Paragraph("Gestión de Personal Camaronero", styles["Normal"]), Paragraph(f"{title} · {start.strftime('%d/%m/%Y')} - {end.strftime('%d/%m/%Y')}", styles["Normal"]), Spacer(1, 14)]
    data = [["Fecha", "Cédula", "Trabajador", "Cargo", "Grupo", "Entrada", "Salida", "Estado", "Observación"]]
    data.extend([[row["fecha"], row["cedula"], f"{row['nombres']} {row['apellidos']}", row["cargo"], row["grupo"], row["hora_entrada"] or "-", row["hora_salida"] or "-", row["estado"], row["observacion"] or "-"] for row in rows])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B1F3A")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#DDE5EF")), ("FONTSIZE", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.extend([table, Spacer(1, 12), Paragraph("Totales: " + " · ".join(f"{key}: {value}" for key, value in summary.items()), styles["Normal"]), Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"])])
    document.build(story)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="aqua_control_reporte.pdf", mimetype="application/pdf")


@app.route("/reportes/exportar/<format_name>")
@role_required("Administrador", "Supervisor", "Consulta")
def exportar_reporte(format_name):
    if format_name not in {"excel", "pdf"}:
        return redirect(url_for("reportes"))
    return export_report(format_name)


@app.route("/usuarios", methods=["GET", "POST"])
@role_required("Administrador")
def usuarios():
    if request.method == "POST":
        user_id = request.form.get("user_id", "").strip()
        name = request.form.get("nombre", "").strip()
        username = request.form.get("usuario", "").strip()
        role = request.form.get("rol", "").strip()
        password = request.form.get("password", "")
        if not name or not username or role not in {"Administrador", "Supervisor", "Consulta"} or (not user_id and not password):
            flash("Completa los datos obligatorios del usuario.", "error")
        else:
            try:
                with get_connection() as connection:
                    if user_id:
                        connection.execute(
                            "UPDATE usuarios SET nombre = ?, usuario = ?, rol = ? WHERE id = ?",
                            (name, username, role, user_id),
                        )
                        if password:
                            connection.execute(
                                "UPDATE usuarios SET password_hash = ? WHERE id = ?",
                                (generate_password_hash(password), user_id),
                            )
                        audit(current_user_id(), "Edición de usuario", username, connection)
                    else:
                        connection.execute(
                            """
                            INSERT INTO usuarios
                            (nombre, usuario, password_hash, rol, estado, fecha_creacion)
                            VALUES (?, ?, ?, ?, 'Activo', ?)
                            """,
                            (name, username, generate_password_hash(password), role, datetime.now().isoformat(timespec="seconds")),
                        )
                        audit(current_user_id(), "Creación de usuario", username, connection)
                flash("Usuario guardado correctamente.", "success")
                return redirect(url_for("usuarios"))
            except sqlite3.IntegrityError:
                flash("Ese nombre de usuario ya existe.", "error")
    edit_id = request.args.get("editar", "").strip()
    with get_connection() as connection:
        user_to_edit = connection.execute("SELECT * FROM usuarios WHERE id = ?", (edit_id,)).fetchone() if edit_id else None
        user_list = connection.execute("SELECT * FROM usuarios ORDER BY nombre").fetchall()
    return render_template("usuarios.html", users=user_list, user=user_to_edit)


@app.post("/usuarios/<int:user_id>/estado")
@role_required("Administrador")
def cambiar_estado_usuario(user_id):
    with get_connection() as connection:
        user = connection.execute("SELECT usuario, estado FROM usuarios WHERE id = ?", (user_id,)).fetchone()
        if user and user_id != current_user_id():
            new_state = "Inactivo" if user["estado"] == "Activo" else "Activo"
            connection.execute("UPDATE usuarios SET estado = ? WHERE id = ?", (new_state, user_id))
            audit(current_user_id(), "Desactivación de usuario" if new_state == "Inactivo" else "Activación de usuario", user["usuario"], connection)
            flash(f"Usuario {new_state.lower()} correctamente.", "success")
    return redirect(url_for("usuarios"))


@app.post("/usuarios/<int:user_id>/restablecer")
@role_required("Administrador")
def restablecer_password_usuario(user_id):
    password = request.form.get("password", "")
    if len(password) < 8:
        flash("La contraseña debe tener al menos 8 caracteres.", "error")
    else:
        with get_connection() as connection:
            connection.execute("UPDATE usuarios SET password_hash = ? WHERE id = ?", (generate_password_hash(password), user_id))
        audit(current_user_id(), "Restablecimiento de contraseña", str(user_id))
        flash("Contraseña restablecida correctamente.", "success")
    return redirect(url_for("usuarios"))


def weekday_name(value):
    if isinstance(value, str):
        value = date.fromisoformat(value)
    return ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"][value.weekday()]


@app.route("/logout")
def logout():
    if current_user_id():
        audit(current_user_id(), "Cierre de sesión")
    session.clear()
    return redirect(url_for("login"))


if __name__ == "__main__":
    app.run(
        host=os.environ.get("FLASK_HOST", "0.0.0.0"),
        port=int(os.environ.get("FLASK_PORT", "5000")),
        debug=app.config["DEBUG"],
    )
